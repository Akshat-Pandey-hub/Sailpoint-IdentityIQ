import getpass
import os
from datetime import datetime, time
import psycopg2
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter


HOST = "localhost"
PORT = 5432
DATABASE = "my_local_db"
USER = "postgres"
SCHEMA = "iiq_native"

TABLES = [
    ("kf_entitlement_assignment", "Entitlement_Assignment.xlsx"),
    ("kf_workgroup_members", "Workgroup_Members.xlsx"),
    ("kf_entitlement_certification", "Entitlement_Certification.xlsx"),
    ("kf_entitlement_certification_status", "Entitlement_Certification_Status.xlsx"),
]


def excel_safe_value(value):
    """
    Convert PostgreSQL timezone-aware datetime/time values into
    Excel-compatible values.

    This changes only the representation in the Excel file.
    PostgreSQL data is never modified.
    """
    if isinstance(value, datetime):
        if value.tzinfo is not None:
            return value.replace(tzinfo=None)
        return value

    if isinstance(value, time):
        if value.tzinfo is not None:
            return value.replace(tzinfo=None)
        return value

    return value


def export_table(cursor, table_name, output_file):
    print()
    print(f"Reading {SCHEMA}.{table_name} ...")

    cursor.execute(
        f'SELECT * FROM "{SCHEMA}"."{table_name}"'
    )

    rows = cursor.fetchall()
    columns = [desc[0] for desc in cursor.description]

    print(f"  Rows    : {len(rows)}")
    print(f"  Columns : {len(columns)}")

    wb = Workbook()
    ws = wb.active
    ws.title = table_name[:31]

    # Header
    for col_num, column_name in enumerate(columns, start=1):
        cell = ws.cell(
            row=1,
            column=col_num,
            value=column_name
        )
        cell.font = Font(bold=True)
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Data
    for row_num, row_data in enumerate(rows, start=2):
        for col_num, value in enumerate(row_data, start=1):

            safe_value = excel_safe_value(value)

            cell = ws.cell(
                row=row_num,
                column=col_num,
                value=safe_value
            )

            if isinstance(safe_value, str) and len(safe_value) > 80:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )
            else:
                cell.alignment = Alignment(
                    vertical="top"
                )

    # Freeze header
    ws.freeze_panes = "A2"

    # Filter
    if columns:
        last_column = get_column_letter(len(columns))
        last_row = len(rows) + 1
        ws.auto_filter.ref = f"A1:{last_column}{last_row}"

    # Column widths
    for col_num, column_name in enumerate(columns, start=1):
        max_length = len(str(column_name))

        for row_data in rows:
            value = row_data[col_num - 1]

            if value is not None:
                value_length = len(str(value))
                max_length = max(max_length, value_length)

        width = min(max(max_length + 2, 12), 50)

        ws.column_dimensions[
            get_column_letter(col_num)
        ].width = width

    ws.row_dimensions[1].height = 22

    wb.save(output_file)

    print(f"  Created : {os.path.abspath(output_file)}")


def main():
    print("=" * 70)
    print("SailPoint IIQ -> KeyForge PostgreSQL -> Excel Export")
    print("=" * 70)

    password = getpass.getpass(
        "Enter PostgreSQL password for user 'postgres': "
    )

    print()
    print("Connecting to PostgreSQL...")
    print(f"Host     : {HOST}")
    print(f"Port     : {PORT}")
    print(f"Database : {DATABASE}")
    print(f"Schema   : {SCHEMA}")
    print(f"User     : {USER}")

    connection = None
    cursor = None

    try:
        connection = psycopg2.connect(
            host=HOST,
            port=PORT,
            dbname=DATABASE,
            user=USER,
            password=password
        )

        cursor = connection.cursor()

        print("Connection successful.")

        output_folder = os.path.join(
            os.getcwd(),
            "sir_share"
        )

        os.makedirs(output_folder, exist_ok=True)

        print()
        print(f"Output folder: {os.path.abspath(output_folder)}")

        for table_name, filename in TABLES:
            output_file = os.path.join(
                output_folder,
                filename
            )

            export_table(
                cursor,
                table_name,
                output_file
            )

        print()
        print("=" * 70)
        print("DONE")
        print("=" * 70)
        print()
        print("Files created in:")
        print(os.path.abspath(output_folder))
        print()
        print("1. Entitlement_Assignment.xlsx")
        print("2. Workgroup_Members.xlsx")
        print("3. Entitlement_Certification.xlsx")
        print("4. Entitlement_Certification_Status.xlsx")
        print()

    except Exception as e:
        print()
        print("=" * 70)
        print("ERROR")
        print("=" * 70)
        print(str(e))
        print()
        print("No PostgreSQL data was modified.")
        print()

    finally:
        if cursor is not None:
            cursor.close()

        if connection is not None:
            connection.close()


if __name__ == "__main__":
    main()